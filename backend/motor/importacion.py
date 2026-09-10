"""Lectura de estados financieros desde Excel, CSV y PDF.

El problema de importar no es abrir el archivo: es que cada empresa nombra sus
cuentas distinto. "Deudores comerciales", "Cartera clientes" y "CxC" son la
misma cosa, y el motor necesita un codigo exacto para calcular.

Por eso la importacion tiene dos etapas separadas a proposito:

  1. LECTURA (este modulo, deterministica): abre el archivo, arma una tabla de
     filas con etiqueta y cifras, detecta los periodos y propone un mapeo con
     un diccionario de sinonimos escrito a mano.
  2. CONFIRMACION (el usuario, en pantalla): lo que el diccionario no reconoce
     se le puede pasar a un modelo para que PROPONGA a que cuenta corresponde
     (ver narrativa.proponer_cuentas). Esa propuesta nunca se aplica sola: se
     muestra marcada, con su confianza y con quien la hizo, y el usuario
     confirma o corrige antes de que se calcule nada.

La regla de la aplicacion se mantiene intacta: el modelo puede opinar sobre
COMO SE LLAMA una fila, jamas sobre cuanto vale ni sobre que significa. Los
numeros salen del archivo y los calculos, del motor.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import asdict, dataclass, field

# Cuentas que entienden los modelos del motor, agrupadas como las espera
# EstadosFinancieros. El orden es el de presentacion en pantalla.
CUENTAS_BALANCE = [
    "efectivo", "cuentas_por_cobrar", "inventarios", "activo_corriente",
    "propiedad_planta_equipo", "activo_total", "proveedores",
    "deuda_financiera_cp", "pasivo_corriente", "deuda_financiera_lp",
    "patrimonio",
]
CUENTAS_RESULTADOS = [
    "ventas", "costo_ventas", "utilidad_bruta", "gastos_operacionales",
    "utilidad_operacional", "gastos_financieros", "utilidad_antes_impuestos",
    "impuestos", "utilidad_neta", "depreciacion", "compras", "ventas_credito",
]
CUENTAS = CUENTAS_BALANCE + CUENTAS_RESULTADOS

GRUPO = {c: "balance" for c in CUENTAS_BALANCE}
GRUPO.update({c: "resultados" for c in CUENTAS_RESULTADOS})

# Diccionario de sinonimos. Se compara sobre el texto normalizado (sin tildes,
# en minusculas). Cada entrada es una frase que debe aparecer completa dentro
# de la etiqueta de la fila. Se evalua de la mas larga a la mas corta, para que
# "utilidad antes de impuestos" gane sobre "utilidad".
SINONIMOS: dict[str, tuple[str, ...]] = {
    "efectivo": ("efectivo y equivalentes", "efectivo", "caja y bancos", "caja",
                 "bancos", "disponible"),
    "cuentas_por_cobrar": ("cuentas comerciales por cobrar",
                           "cuentas comerciales y otras cuentas por cobrar",
                           "deudores comerciales y otras cuentas por cobrar","cuentas por cobrar", "deudores comerciales", "deudores",
                           "cartera clientes", "cartera", "clientes", "cxc"),
    "inventarios": ("inventarios", "inventario", "existencias", "mercancias",
                    "mercancia"),
    "activo_corriente": ("total activo corriente", "activo corriente",
                         "activos corrientes", "total corriente activo"),
    "propiedad_planta_equipo": ("propiedad planta y equipo", "propiedades planta y equipo",
                                "propiedad planta", "activos fijos", "activo fijo",
                                "ppe", "planta y equipo"),
    "activo_total": ("total activo", "total activos", "activo total", "activos totales",
                     "suma del activo"),
    "proveedores": ("proveedores", "cuentas por pagar proveedores",
                    "acreedores comerciales", "cuentas por pagar comerciales",
                    "cuentas comerciales por pagar",
                    "cuentas comerciales y otras cuentas por pagar",
                    "cuentas por pagar y otros pasivos"),
    "deuda_financiera_cp": ("obligaciones financieras corto plazo",
                            "obligaciones financieras cp", "deuda financiera corto plazo",
                            "deuda financiera cp", "creditos corto plazo",
                            "obligaciones bancarias corto plazo"),
    "pasivo_corriente": ("total pasivo corriente", "pasivo corriente",
                         "pasivos corrientes", "total corriente pasivo"),
    "deuda_financiera_lp": ("obligaciones financieras largo plazo",
                            "obligaciones financieras lp", "deuda financiera largo plazo",
                            "deuda financiera lp", "pasivo no corriente",
                            "pasivos no corrientes", "creditos largo plazo",
                            "obligaciones bancarias largo plazo"),
    "patrimonio": ("total patrimonio", "patrimonio neto", "patrimonio",
                   "capital contable"),
    "ventas": ("ingresos operacionales", "ingresos por ventas", "ingresos de actividades",
               "ventas netas", "ventas", "ingresos"),
    "costo_ventas": ("costo de ventas", "costo de la mercancia vendida",
                     "costo de mercancia vendida", "costos de ventas", "costo ventas"),
    "utilidad_bruta": ("utilidad bruta", "ganancia bruta", "margen bruto en pesos"),
    "gastos_operacionales": ("gastos de administracion y ventas", "gastos operacionales",
                             "gastos de operacion", "gastos administrativos y de ventas",
                             "gastos administracion y ventas"),
    "utilidad_operacional": ("utilidad operacional", "utilidad de operacion",
                             "resultado operacional", "ganancia operacional",
                             "utilidad antes de intereses e impuestos", "uaii", "ebit"),
    "gastos_financieros": ("gastos financieros", "intereses pagados",
                           "costos financieros", "gasto por intereses", "intereses"),
    "utilidad_antes_impuestos": ("utilidad antes de impuestos", "utilidad antes de impuesto",
                                 "resultado antes de impuestos", "uai",
                                 "ganancia antes de impuestos"),
    "impuestos": ("impuesto de renta", "impuesto sobre la renta", "provision de impuestos",
                  "impuestos", "impuesto"),
    "utilidad_neta": ("utilidad neta", "ganancia neta", "resultado del ejercicio",
                      "resultado neto", "utilidad del ejercicio"),
    "depreciacion": ("depreciacion y amortizacion", "depreciacion", "amortizacion"),
    "compras": ("compras del periodo", "compras"),
    "ventas_credito": ("ventas a credito", "ventas credito"),
}

# Filas que nunca son una cuenta: encabezados, totales de control, notas.
RUIDO = ("cuenta", "concepto", "descripcion", "rubro", "nota", "notas",
         "estado de situacion", "estado de resultados", "balance general",
         "cifras en", "expresado en", "variacion", "analisis", "%")

MAX_PERIODOS = 6


@dataclass
class Fila:
    """Una linea del archivo: su etiqueta original y las cifras que traia."""

    etiqueta: str
    valores: list
    origen: str = ""            # hoja o pagina de donde salio
    cuenta: str | None = None   # codigo asignado
    grupo: str = ""             # balance o resultados
    asignado_por: str = ""      # "diccionario", "modelo" o "usuario"
    confianza: str = ""         # alta, media o baja
    nota: str = ""


@dataclass
class Tabla:
    """Lo que se logro leer del archivo, antes de que nadie confirme nada."""

    periodos: list[str] = field(default_factory=list)
    filas: list[Fila] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    origen: str = ""

    def como_dict(self) -> dict:
        return {
            "periodos": self.periodos,
            "filas": [asdict(f) for f in self.filas],
            "avisos": self.avisos,
            "origen": self.origen,
        }


# ------------------------------------------------------------ texto y numeros


def normalizar(texto: str) -> str:
    """Minusculas, sin tildes y sin puntuacion, para poder comparar etiquetas."""
    if texto is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = t.lower().replace("$", " ")
    t = re.sub(r"[^\w\s%]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def leer_numero(bruto):
    """Interpreta una celda como cifra, o devuelve None si no lo es.

    Tiene que aguantar lo que sale de una contabilidad colombiana:

        "1.234.567,89"   -> 1234567.89   (punto de miles, coma decimal)
        "1,234,567.89"   -> 1234567.89   (formato ingles)
        "$ 2.450"        -> 2450.0
        "(550)"          -> -550.0       (parentesis = negativo)
        "1.234"          -> 1234.0       (miles, no un decimal)
        "-", "", "n/d"   -> None         (dato ausente, NO cero)

    Un dato ausente jamas se convierte en cero: el motor distingue "no
    informado" de "vale cero" y todo el diagnostico depende de eso.
    """
    if bruto is None:
        return None
    if isinstance(bruto, bool):
        return None
    if isinstance(bruto, (int, float)):
        return float(bruto)

    t = str(bruto).strip()
    if not t:
        return None

    negativo = t.startswith("(") and t.endswith(")")
    t = t.strip("()").replace("$", "").replace("COP", "").replace("%", "")
    t = t.replace(" ", " ").replace(" ", "")
    if not t or t in {"-", "--", "n/d", "nd", "na", "n/a", "."}:
        return None
    if t.startswith("-"):
        negativo = True
        t = t[1:]
    if not re.fullmatch(r"[\d.,]+", t):
        return None

    coma, punto = t.rfind(","), t.rfind(".")
    if coma >= 0 and punto >= 0:
        # El separador decimal es el que aparece de ultimo.
        if coma > punto:
            t = t.replace(".", "").replace(",", ".")
        else:
            t = t.replace(",", "")
    elif coma >= 0:
        # Una sola coma: decimal si deja uno o dos digitos ("2,5"), miles si no.
        t = t.replace(",", "." if len(t) - coma - 1 <= 2 else "")
    elif punto >= 0:
        # Un solo punto con exactamente tres digitos detras es miles: "1.234".
        if len(t) - punto - 1 == 3:
            t = t.replace(".", "")
    try:
        valor = float(t)
    except ValueError:
        return None
    return -valor if negativo else valor


def parece_periodo(texto) -> str | None:
    """Reconoce un encabezado de columna que nombre un periodo.

    Acepta un ano suelto (2024), un ano con adornos (FY2024, Dic-2024) y
    rangos cortos. Devuelve el periodo ya limpio, o None.
    """
    if texto is None:
        return None
    t = str(texto).strip()
    if not t:
        return None
    anios = re.findall(r"(19|20)\d{2}", t)
    if not anios:
        return None
    completo = re.findall(r"(?:19|20)\d{2}", t)
    if len(completo) == 1 and len(t) <= 24:
        return t if len(t) <= 12 else completo[0]
    return completo[-1]


# --------------------------------------------------------------- diccionario


def _frases_ordenadas():
    """Sinonimos de la frase mas larga a la mas corta, con su cuenta."""
    pares = [(frase, cuenta) for cuenta, frases in SINONIMOS.items() for frase in frases]
    return sorted(pares, key=lambda p: -len(p[0]))


FRASES = _frases_ordenadas()


def buscar_cuenta(etiqueta: str):
    """Cuenta que corresponde a una etiqueta segun el diccionario.

    Devuelve (codigo, confianza) o (None, ""). La confianza es alta cuando la
    etiqueta es practicamente el nombre de la cuenta, y media cuando el
    sinonimo aparece dentro de un texto mas largo, donde puede haber trampa
    ("total activo corriente" contiene "activo corriente" pero tambien
    "corriente" a secas).
    """
    limpia = normalizar(etiqueta)
    if not limpia or limpia in RUIDO:
        return None, ""
    for frase, cuenta in FRASES:
        if limpia == frase:
            return cuenta, "alta"
    for frase, cuenta in FRASES:
        if re.search(r"\b" + re.escape(frase) + r"\b", limpia):
            sobra = len(limpia) - len(frase)
            return cuenta, "alta" if sobra <= 6 else "media"
    return None, ""


def mapear(tabla: Tabla) -> Tabla:
    """Asigna cuentas a las filas con el diccionario, sin tocar los numeros.

    Si dos filas reclaman la misma cuenta se conserva la de mejor confianza y
    la otra queda libre, con una nota que lo explica. Adivinar en silencio ahi
    seria la forma mas facil de arruinar un balance.
    """
    tomadas: dict[str, Fila] = {}
    for fila in tabla.filas:
        cuenta, confianza = buscar_cuenta(fila.etiqueta)
        if not cuenta:
            continue
        previa = tomadas.get(cuenta)
        if previa is None:
            fila.cuenta, fila.confianza, fila.asignado_por = cuenta, confianza, "diccionario"
            fila.grupo = GRUPO[cuenta]
            tomadas[cuenta] = fila
            continue
        # Ya habia una fila para esa cuenta: gana la de confianza alta.
        if confianza == "alta" and previa.confianza != "alta":
            previa.cuenta, previa.grupo, previa.asignado_por = None, "", ""
            previa.nota = (f"Otra fila se identifico mejor como {cuenta}. "
                           "Reviselo si esta fila era la correcta.")
            previa.confianza = ""
            fila.cuenta, fila.confianza, fila.asignado_por = cuenta, confianza, "diccionario"
            fila.grupo = GRUPO[cuenta]
            tomadas[cuenta] = fila
        else:
            fila.nota = (f"Se parece a {cuenta}, pero esa cuenta ya la tomo la fila "
                         f"\"{previa.etiqueta}\". Elija cual de las dos es.")
    return tabla


# -------------------------------------------------------------------- lectura


def _filas_desde_matriz(matriz, origen: str, tabla: Tabla,
                        fusionar_continuaciones: bool = False) -> None:
    """Convierte una matriz de celdas en filas con etiqueta y cifras.

    La primera columna que trae texto es la etiqueta; las columnas numericas
    son los periodos, en el orden en que aparecen. Se detectan los encabezados
    de periodo en las primeras filas.
    """
    if not matriz:
        return

    # Encabezado: la primera fila donde aparezcan periodos y TODA cifra sea un
    # ano. Un "2024" se lee como numero, asi que no sirve exigir que la fila no
    # tenga cifras: hay que exigir que las que tenga sean todas periodos. Si la
    # etiqueta de la fila ya es una cuenta conocida, no es encabezado sino
    # datos (una fila "Ventas | 2023 | 2024" no puede pasar por titulo).
    encabezado, periodos = None, []
    for i, fila in enumerate(matriz[:8]):
        limpios = [p for p in (parece_periodo(c) for c in fila) if p]
        if not limpios:
            continue
        numericas = [c for c in fila if leer_numero(c) is not None]
        if not all(parece_periodo(c) for c in numericas):
            continue
        etiqueta = next((c for c in fila if c is not None and str(c).strip()
                         and leer_numero(c) is None), "")
        if etiqueta and buscar_cuenta(str(etiqueta))[0]:
            continue
        encabezado, periodos = i, limpios[:MAX_PERIODOS]
        break

    if periodos and not tabla.periodos:
        tabla.periodos = periodos

    inicio = 0 if encabezado is None else encabezado + 1
    for fila in matriz[inicio:]:
        # Una celda puede traer su texto partido en varias lineas: se une.
        celdas = [None if c is None else str(c).replace("\n", " ").strip()
                  if isinstance(c, str) else c for c in fila]
        etiqueta = next((str(c).strip() for c in celdas
                         if c is not None and str(c).strip()
                         and leer_numero(c) is None), "")
        if not etiqueta:
            continue
        numeros = [leer_numero(c) for c in celdas]
        numeros = [n for n in numeros if n is not None]
        if not numeros:
            # Renglon sin cifras. En un PDF suele ser la segunda linea de un
            # nombre largo ("Efectivo y" / "equivalentes"), que el extractor
            # parte en dos. Se pega a la fila anterior en vez de perderse.
            if (fusionar_continuaciones and tabla.filas
                    and len(etiqueta.split()) <= 4
                    and not buscar_cuenta(etiqueta)[0]):
                anterior = tabla.filas[-1]
                if anterior.origen == origen:
                    anterior.etiqueta = f"{anterior.etiqueta} {etiqueta}".strip()
            continue
        tabla.filas.append(Fila(etiqueta=etiqueta,
                                valores=numeros[:MAX_PERIODOS],
                                origen=origen))


def leer_excel(contenido: bytes, tabla: Tabla) -> None:
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    for hoja in libro.worksheets:
        matriz = [list(f) for f in hoja.iter_rows(values_only=True)]
        _filas_desde_matriz(matriz, hoja.title, tabla)
    libro.close()
    if len(libro.sheetnames) > 1:
        tabla.avisos.append(
            f"El archivo trae {len(libro.sheetnames)} hojas y se leyeron todas. "
            "Revise que no se hayan mezclado cuentas de hojas distintas.")


def leer_csv(contenido: bytes, tabla: Tabla) -> None:
    texto = None
    for codificacion in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        raise ValueError("No se pudo leer el archivo de texto.")
    try:
        dialecto = csv.Sniffer().sniff(texto[:4000], delimiters=";,\t|")
        delimitador = dialecto.delimiter
    except csv.Error:
        delimitador = ";" if texto.count(";") > texto.count(",") else ","
    matriz = [f for f in csv.reader(io.StringIO(texto), delimiter=delimitador)]
    _filas_desde_matriz(matriz, "csv", tabla)


# Una linea de tabla de contenido: texto, una fila de puntos y el numero de
# pagina. El numero NO es una cifra contable, pero se parece a una, y por eso
# un informe anual completo entraba con los numeros de pagina como si fueran
# saldos. Verificado con el informe de Ecopetrol 2024, cuyo indice produjo
# "Propiedades, planta y equipo = 63".
_INDICE = re.compile(r"\.{4,}\s*\d{1,4}\s*$")


def _es_linea_de_indice(linea: str) -> bool:
    return bool(_INDICE.search(linea.strip()))


def _filas_utiles(filas) -> int:
    """Cuantas de estas filas traen al menos una cifra.

    Es el criterio para escoger entre leer una pagina como tabla o como texto:
    gana la via que rescate mas renglones con numeros.
    """
    return sum(1 for f in filas if any(v is not None for v in f.valores))


def _paginas_pedidas(rango: str, total: int) -> set[int] | None:
    """Traduce "10-14" o "10,11,12" a un conjunto de paginas, 1-based.

    Devuelve None si no se pidio nada, que significa "todas".
    """
    if not rango or not rango.strip():
        return None
    elegidas = set()
    for parte in rango.replace(" ", "").split(","):
        if not parte:
            continue
        if "-" in parte:
            a, _, b = parte.partition("-")
            try:
                ini, fin = int(a), int(b)
            except ValueError:
                raise ValueError(f"Rango de paginas invalido: '{parte}'.")
            if ini > fin:
                ini, fin = fin, ini
            elegidas.update(range(max(1, ini), min(total, fin) + 1))
        else:
            try:
                n = int(parte)
            except ValueError:
                raise ValueError(f"Numero de pagina invalido: '{parte}'.")
            if 1 <= n <= total:
                elegidas.add(n)
    if not elegidas:
        raise ValueError(
            f"Ninguna de las paginas pedidas existe. El PDF tiene {total}.")
    return elegidas


def leer_pdf(contenido: bytes, tabla: Tabla, paginas: str = "") -> None:
    """Extrae las cifras de un PDF, decidiendo pagina por pagina como leerla.

    Un PDF no tiene celdas: la tabla hay que reconstruirla. Por eso esta via
    siempre deja un aviso pidiendo revision, y por eso ninguna cifra entra al
    calculo sin que el usuario la confirme en pantalla.

    Tres defensas que salieron de probar informes anuales reales:

    1. **La eleccion tabla-o-texto es por pagina, no global.** Antes bastaba
       que UNA pagina reportara una tabla para que todas las demas dejaran de
       leerse como texto. En el informe de Nutresa pdfplumber detecta siete
       "tablas" espurias en la pagina del balance, y eso hacia que el balance
       -que ahi es texto perfectamente legible- se perdiera entero.

    2. **Las lineas de tabla de contenido se descartan.** Un indice tiene la
       misma forma que un estado financiero -concepto a la izquierda, numero a
       la derecha- pero ese numero es una pagina. El informe de Ecopetrol
       producia "Propiedades, planta y equipo = 63".

    3. **Se cuentan las paginas que son imagen sin texto.** Los estados
       firmados suelen ir escaneados, y entonces lo unico legible del
       documento son el indice y las notas: se leeria todo menos lo que
       importa.
    """
    import pdfplumber

    paginas_por_tabla = 0
    paginas_por_texto = 0
    lineas_indice = 0
    paginas_imagen = 0
    paginas_con_texto = 0

    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        if not pdf.pages:
            raise ValueError("El PDF no tiene paginas legibles.")

        total_paginas = len(pdf.pages)
        pedidas = _paginas_pedidas(paginas, total_paginas)
        deteccion = None

        # Si nadie dijo que paginas leer, se buscan solas. Un informe anual
        # entero mezcla las tablas de las notas con el balance, y pedirle al
        # usuario que abra el PDF y anote las paginas es un trabajo que la
        # maquina puede hacer.
        if pedidas is None and total_paginas > 8:
            from .secciones import analizar_paginas
            textos = [(pag.extract_text() or "") for pag in pdf.pages]
            deteccion = analizar_paginas(textos)
            if deteccion["encontrado"]:
                pedidas = _paginas_pedidas(deteccion["paginas"], total_paginas)
                tabla.avisos.append(
                    "Las paginas se detectaron solas. " + deteccion["lectura"] +
                    " Si no es lo que esperaba, indique el rango a mano y vuelva "
                    "a cargar el archivo.")

        for n, pagina in enumerate(pdf.pages, start=1):
            texto = pagina.extract_text() or ""

            # El conteo de escaneadas mira el documento COMPLETO, aunque solo se
            # vayan a leer unas paginas. Si no, acotar el rango escondia la
            # advertencia mas importante: que los estados van escaneados y que
            # lo unico legible del documento son las notas.
            if texto.strip():
                paginas_con_texto += 1
            elif pagina.images:
                paginas_imagen += 1

            if pedidas is not None and n not in pedidas:
                continue

            # Via A: las tablas que pdfplumber alcance a reconocer.
            porA = Tabla(periodos=list(tabla.periodos))
            for matriz in pagina.extract_tables() or []:
                if matriz and len(matriz) > 1:
                    _filas_desde_matriz(matriz, f"pagina {n}", porA)

            # Via B: el texto plano, renglon por renglon, sin el indice.
            utiles = []
            for linea in texto.splitlines():
                if not linea.strip():
                    continue
                if _es_linea_de_indice(linea):
                    lineas_indice += 1
                    continue
                utiles.append(linea)
            porB = Tabla(periodos=list(tabla.periodos))
            matriz = [m for m in (_partir_linea(l) for l in utiles) if m]
            _filas_desde_matriz(matriz, f"pagina {n} (texto)", porB,
                                fusionar_continuaciones=True)

            # Gana la que rescate mas renglones con cifras.
            if _filas_utiles(porA.filas) >= _filas_utiles(porB.filas) and porA.filas:
                elegida, cual = porA, "tabla"
            else:
                elegida, cual = porB, "texto"

            if not elegida.filas:
                continue
            if cual == "tabla":
                paginas_por_tabla += 1
            else:
                paginas_por_texto += 1

            tabla.filas.extend(elegida.filas)
            for per in elegida.periodos:
                if per not in tabla.periodos:
                    tabla.periodos.append(per)

    if pedidas is not None and deteccion is None:
        tabla.avisos.append(
            f"Se leyeron unicamente las paginas pedidas ({len(pedidas)} de "
            f"{total_paginas}). El resto del documento se ignoro.")

    if paginas_por_tabla and paginas_por_texto:
        tabla.avisos.append(
            f"Se leyeron {paginas_por_tabla} paginas como tabla y "
            f"{paginas_por_texto} como texto suelto, segun lo que rescataba mas "
            f"cifras en cada una. Compare cada renglon contra el documento "
            f"original antes de confirmar.")
    elif paginas_por_tabla:
        tabla.avisos.append(
            "Las cifras se reconstruyeron de las tablas del PDF. Un PDF no guarda "
            "celdas, asi que compare cada renglon contra el documento original "
            "antes de confirmar.")
    else:
        tabla.avisos.append(
            "El PDF no traia tablas reconocibles y se leyo como texto suelto, que "
            "es la via mas fragil. Revise renglon por renglon.")

    if lineas_indice:
        tabla.avisos.append(
            f"Se descartaron {lineas_indice} lineas de tabla de contenido. En un "
            f"indice el numero de la derecha es una pagina, no un saldo.")

    if paginas_imagen and deteccion is not None and deteccion.get("encontrado"):
        tabla.avisos.append(
            "CUIDADO: el rango se eligio solo, pero este documento tiene paginas "
            "escaneadas. Si los estados financieros son justamente esas paginas, "
            "lo que se detecto es un anexo o una nota que se les parece. Abra el "
            "PDF y confirme que lo leido es el balance y el estado de resultados "
            "de verdad.")

    if paginas_imagen:
        total = paginas_imagen + paginas_con_texto
        tabla.avisos.append(
            f"ATENCION: {paginas_imagen} de {total} paginas son imagenes sin texto "
            f"y no se pudieron leer. En los informes anuales los estados firmados "
            f"suelen ir escaneados, asi que es probable que lo que se alcanzo a "
            f"leer sean las notas y no los estados. Verifique que las cifras de "
            f"abajo salgan del balance y del estado de resultados; si no, "
            f"consiga el archivo en Excel o digite las cuentas a mano.")


def _partir_linea(linea: str):
    """Separa una linea de texto de PDF en etiqueta y cifras.

    Las cifras van pegadas al final del renglon separadas por espacios; el
    resto es el nombre de la cuenta.
    """
    texto = linea.strip()
    piezas = re.split(r"\s{2,}|\t", texto)
    if len(piezas) > 1:
        return [p.strip() for p in piezas if p.strip()]

    # Separado por espacios sueltos. Las cifras son la cola del renglon, asi
    # que se corta ahi: partir por los ultimos espacios rompia los nombres
    # largos ("Efectivo y equivalentes 450 380" perdia "equivalentes").
    #
    # La cola admite el signo de moneda delante de cada cifra, porque asi
    # vienen casi todos los estados financieros publicados en Colombia:
    # "Total activo corriente $ 8.817.990 $ 6.416.225". Sin esto el "$"
    # intermedio cortaba la cola y solo se rescataba la ultima columna.
    cola = re.search(r"^(.*?)((?:\s+\$?\s*\(?-?[\d.,]+\)?%?)+\s*\$?)$", texto)
    if not cola:
        return [texto]
    etiqueta = cola.group(1).strip()
    # El "$" es ruido tipografico, no un dato: se descarta al separar.
    cifras = [c for c in cola.group(2).replace("$", " ").split() if c]
    return ([etiqueta] if etiqueta else []) + cifras


EXTENSIONES = {
    ".xlsx": leer_excel, ".xlsm": leer_excel,
    ".csv": leer_csv, ".txt": leer_csv,
    ".pdf": leer_pdf,
}


def importar(nombre: str, contenido: bytes, paginas: str = "") -> Tabla:
    """Lee un archivo y devuelve la tabla con el mapeo ya propuesto.

    No calcula ni valida nada financiero: eso es del motor, y solo despues de
    que el usuario confirme el mapeo en pantalla.

    `paginas` solo aplica a PDF y acota que paginas se leen ("10-14" o "10,11").
    Un informe anual trae los estados en dos o tres paginas y ochenta de notas;
    leerlo entero mezcla las tablas de las notas con el balance.
    """
    extension = "." + nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""
    lector = EXTENSIONES.get(extension)
    if lector is None:
        raise ValueError(
            f"Formato no soportado ({extension or 'sin extension'}). "
            "Se aceptan Excel (.xlsx), CSV (.csv) y PDF (.pdf).")
    if not contenido:
        raise ValueError("El archivo llego vacio.")

    tabla = Tabla(origen=nombre)
    if extension == ".pdf":
        lector(contenido, tabla, paginas)
    else:
        lector(contenido, tabla)

    if not tabla.filas:
        raise ValueError(
            "No se encontro ninguna fila con cifras. Revise que el archivo tenga "
            "una columna con los nombres de las cuentas y al menos una con numeros.")

    if not tabla.periodos:
        columnas = max(len(f.valores) for f in tabla.filas)
        tabla.periodos = [f"Periodo {i + 1}" for i in range(columnas)]
        tabla.avisos.append(
            "No se reconocieron los anos en los encabezados, asi que las columnas "
            "quedaron como Periodo 1, 2... Renombrelas antes de analizar.")

    mapear(tabla)

    sin_asignar = [f for f in tabla.filas if not f.cuenta]
    if sin_asignar:
        tabla.avisos.append(
            f"{len(sin_asignar)} de {len(tabla.filas)} filas no se reconocieron con "
            "el diccionario. Puede asignarlas a mano o pedirle una propuesta al "
            "modelo, que igual tendra que confirmar.")
    return tabla


# ------------------------------------------------------- armado del resultado


def armar_estados(tabla_cruda: dict, empresa: str = "", moneda: str = "COP",
                  unidad: str = "millones") -> dict:
    """Convierte la tabla ya confirmada en el JSON que come el motor.

    Recibe el mapeo tal como quedo en pantalla, con las correcciones del
    usuario. Las cuentas que nadie asigno simplemente no van: el motor sabe
    trabajar con cuentas ausentes y reportarlas como no disponibles.
    """
    periodos = [str(p) for p in tabla_cruda.get("periodos") or []]
    if not periodos:
        raise ValueError("La tabla no tiene periodos.")

    datos = {
        "empresa": empresa or tabla_cruda.get("empresa") or "Empresa importada",
        "moneda": moneda,
        "unidad": unidad,
        "periodos": periodos,
        "balance": {},
        "resultados": {},
    }

    # Cuentas que por naturaleza son un monto positivo. Los estados
    # financieros publicados las presentan entre parentesis, es decir en
    # negativo, porque los estan restando dentro del encadenamiento:
    # "Costos de ventas (12.457.813)". El motor las espera positivas.
    #
    # Se voltean, pero NUNCA en silencio: cada volteo queda anotado en los
    # supuestos y se reporta. Verificado con el informe de Grupo Nutresa.
    SIEMPRE_POSITIVAS = ("costo_ventas", "gastos_operacionales",
                         "gastos_financieros", "impuestos", "compras",
                         "depreciacion")
    ajustes = []

    usadas = set()
    for fila in tabla_cruda.get("filas") or []:
        cuenta = fila.get("cuenta")
        if not cuenta or cuenta not in GRUPO or cuenta in usadas:
            continue
        valores = [leer_numero(v) for v in (fila.get("valores") or [])]
        # Todas las cuentas tienen que traer un valor por periodo: si faltan se
        # completan con None, nunca con cero.
        #
        # Cuando SOBRAN se conservan los ultimos, no los primeros. En un estado
        # financiero publicado la columna de la izquierda es el numero de nota:
        # "Inventarios | 11 | 2.558.764 | 2.447.873". Recortando por la
        # izquierda se guardaba el 11 como si fuera el saldo del primer
        # periodo. Verificado con el informe consolidado de Grupo Nutresa.
        n = len(periodos)
        if len(valores) > n:
            valores = valores[-n:]
        else:
            valores = valores + [None] * (n - len(valores))
        if all(v is None for v in valores):
            continue

        if cuenta in SIEMPRE_POSITIVAS and any(v is not None and v < 0 for v in valores):
            valores = [None if v is None else abs(v) for v in valores]
            ajustes.append(
                f"'{cuenta}' venia en negativo y se volteo a positivo. Los "
                f"estados publicados presentan costos y gastos entre parentesis "
                f"porque los estan restando; el motor los espera como monto.")

        datos[GRUPO[cuenta]][cuenta] = valores
        usadas.add(cuenta)

    if not datos["balance"] and not datos["resultados"]:
        raise ValueError(
            "Ninguna fila quedo asignada a una cuenta conocida. Asigne al menos "
            "las cuentas principales antes de analizar.")

    if ajustes:
        datos["supuestos"] = {"ajustes_importacion": ajustes}
    return datos
